"""Read exports incrementally and check counts and Markdown references."""

import csv
import json
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from mailanalyst.exports.json_stream import records
from mailanalyst.cancellation import check_cancel

csv.field_size_limit(min(sys.maxsize, 2**31 - 1))


def validate_output(path: Path, count: int, cancel=None) -> None:
    check_cancel(cancel)
    suffix = path.suffix.lower()
    if suffix == ".json":
        with path.open(encoding="utf-8") as file:
            actual = sum(1 for _ in checked(records(file), cancel))
    elif suffix == ".parquet":
        import pyarrow.parquet as pq
        with pq.ParquetFile(path) as file:
            actual = sum(batch.num_rows for batch in checked(file.iter_batches(batch_size=500), cancel))
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True)
        try:
            actual = max(0, sum(1 for _ in checked(book.active.iter_rows(), cancel)) - 1)
        finally:
            book.close()
    elif suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as file:
            actual = max(0, sum(1 for _ in checked(csv.reader(file), cancel)) - 1)
    elif suffix == ".xml":
        actual = 0
        with path.open("rb") as file:
            events = ET.iterparse(file, events=("start", "end"))
            _, root = next(events)
            for event, element in checked(events, cancel):
                if event == "end" and element.tag == "email":
                    actual += 1
                    root.clear()
    elif suffix in {".md", ".markdown"}:
        with path.open(encoding="utf-8") as file:
            header = "".join(file.readline() for _ in range(3))
        if header != f"# MailAnalyst Export\n\n{count} Nachrichten\n":
            raise ValueError(f"Invalid Markdown export: {path}")
        return
    else:
        raise ValueError(f"Unsupported validation format: {suffix}")
    check_cancel(cancel)
    if actual != count:
        raise ValueError(f"Export count mismatch: {path}: {actual} != {count}")


def validate_dataset(directory: Path, count: int, cancel=None) -> None:
    validate_output(directory / "index.csv", count, cancel)
    actual, current, markdown = 0, None, None
    try:
        with (directory / "index.jsonl").open(encoding="utf-8") as index:
            for line in checked(index, cancel):
                if not line.strip():
                    continue
                row = json.loads(line)
                path = (directory / row["markdown_file"]).resolve()
                if not path.is_relative_to(directory.resolve()):
                    raise ValueError("Markdown reference leaves dataset")
                if path != current:
                    if markdown:
                        markdown.close()
                    markdown, current = path.open(encoding="utf-8"), path
                anchor = f'<a id="{row["anchor"]}"></a>'
                if not any(text.strip() == anchor for text in checked(markdown, cancel)):
                    raise ValueError("Missing Markdown anchor")
                actual += 1
    finally:
        if markdown:
            markdown.close()
    check_cancel(cancel)
    if actual != count:
        raise ValueError("Markdown index count mismatch")


def checked(iterator, cancel):
    for item in iterator:
        check_cancel(cancel)
        yield item
