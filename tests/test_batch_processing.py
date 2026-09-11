"""Real EML batch/cache behavior and bounded result/export regression tests."""

from contextlib import closing
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
import pyarrow.parquet as pq

from mailanalyst.batch_pipeline import build_store
from mailanalyst.cancellation import Cancellation, Cancelled
from mailanalyst.exports.dispatch import write_output
from mailanalyst.exports.markdown import write_markdown_dataset
from mailanalyst.exports.validation import validate_dataset
from mailanalyst.exports.list_view import ListView
from mailanalyst.pipeline import build_dataframe
from mailanalyst.record_store import RecordStore
from tests.samples import create_sources


class BatchTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source = create_sources(self.root)
        self.cache = self.root / "cache.sqlite3"
        self.number = 0

    def build(self, **kwargs):
        self.number += 1
        store = build_store(self.source, self.cache, self.root / f"work-{self.number}.sqlite3", **kwargs)
        self.addCleanup(store.close)
        return store

    def test_real_eml_cache_migration_invalidation_and_corruption(self):
        build_dataframe(self.source, self.cache)
        first = self.build(workers=2, batch_size=1)
        self.assertEqual([a["mode"] for a in first.audits()], ["parsed", "parsed"])
        cached = self.build(hash_check=True)
        self.assertEqual([a["mode"] for a in cached.audits()], ["cache", "cache"])
        self.assertEqual(list(first.records()), list(cached.records()))
        with (self.source / "first.eml").open("ab") as file:
            file.write(b"\nchanged")
        changed = self.build()
        self.assertEqual([a["mode"] for a in changed.audits()], ["parsed", "cache"])
        with closing(sqlite3.connect(self.cache)) as db:
            db.execute("UPDATE messages SET payload='{}' WHERE seq=2")
            db.commit()
        repaired = self.build()
        self.assertEqual([a["mode"] for a in repaired.audits()], ["cache", "parsed"])
        before = self.cache.read_bytes()
        token = Cancellation()
        with self.assertRaises(Cancelled):
            self.build(cancel=token, progress_callback=lambda *_: token.request())
        self.assertEqual(self.cache.read_bytes(), before)

    def test_stream_exports_match_materialized_sample(self):
        store = self.build(batch_size=1)
        frame = pd.DataFrame(list(store.records()))
        for suffix in ("json", "csv", "xlsx", "parquet", "xml", "md"):
            path = self.root / f"stream.{suffix}"
            write_output(store, path)
            if suffix == "json":
                self.assertEqual(json.loads(path.read_text(encoding="utf-8")), list(store.records()))
            elif suffix == "parquet":
                self.assertEqual(pq.read_table(path).to_pylist(), list(store.records()))
            elif suffix == "xlsx":
                self.assertEqual(pd.read_excel(path)["message_id"].tolist(), frame.message_id.tolist())
            elif suffix in {"csv", "md"}:
                original = self.root / f"original.{suffix}"
                write_output(frame, original)
                self.assertEqual(path.read_bytes(), original.read_bytes())
        write_output(ListView(store), self.root / "review.csv")
        self.assertIn("Betreff", pd.read_csv(self.root / "review.csv").columns)
        for name, data in (("old", frame), ("new", store)):
            write_markdown_dataset(data, self.root / name)
            validate_dataset(self.root / name, 2)
        for path in (self.root / "old").rglob("*.md"):
            self.assertEqual(path.read_bytes(), (self.root / "new" / path.relative_to(self.root / "old")).read_bytes())

    def test_batch_limits_late_types_and_preview_totals(self):
        store = RecordStore(self.root / "records.sqlite3")
        self.addCleanup(store.close)
        store.batch_size = 37
        store.append("synthetic", ({"source_path": f"synthetic::{i}", "source_file_path": "synthetic",
                                    "sent_year": 2026 if i < 1000 else "", "subject": f"={i}",
                                    "body_text_clean": "synthetic " * 10,
                                    "parse_status": "error" if i == 1000 else "ok",
                                    "parse_error": "synthetic error" if i == 1000 else ""}
                                   for i in range(1001)))
        self.assertEqual(store.max_batch_rows, 37)
        preview = store.preview()
        self.assertEqual(len(preview), 500)
        self.assertEqual(preview.attrs["total_messages"], 1001)
        self.assertEqual(preview.attrs["total_errors"], 1)
        path = self.root / "mixed.parquet"
        write_output(store, path)
        result = pq.ParquetFile(path)
        self.assertEqual(result.metadata.num_rows, 1001)
        self.assertLessEqual(max(result.metadata.row_group(i).num_rows for i in range(result.num_row_groups)), 37)
        self.assertEqual(result.read(columns=["sent_year"]).to_pylist()[-1], {"sent_year": ""})
        self.assertEqual(result.read(columns=["subject"]).to_pylist()[0], {"subject": "=0"})
        with patch("mailanalyst.record_store.BATCH_BYTES", 100):
            small = list(store.batches())
            self.assertTrue(all(len(batch) == 1 for batch in small))
            self.assertEqual(len(store.preview()), 1)

    def test_cancel_during_export_preserves_existing_output(self):
        token = Cancellation()
        store = self.build(cancel=token)
        path = self.root / "existing.json"
        path.write_text("old", encoding="utf-8")
        token.request()
        with self.assertRaises(Cancelled):
            write_output(store, path)
        self.assertEqual(path.read_text(encoding="utf-8"), "old")
        self.assertEqual(list(self.root.glob(".pending-*")), [])

    def test_streaming_view_formats_preserve_formula_protection(self):
        from openpyxl import load_workbook
        import csv
        store = RecordStore(self.root / "view.sqlite3")
        self.addCleanup(store.close)
        text = "=HYPERLINK(\"https://example.test\")"
        store.append("synthetic", [{"source_path": "synthetic::1", "source_file_path": "synthetic",
                                    "parse_status": "ok", "subject": text, "number": -7}])
        write_output(store, self.root / "view.csv")
        with (self.root / "view.csv").open(encoding="utf-8-sig", newline="") as file:
            row = list(csv.DictReader(file))[0]
        self.assertEqual(row["subject"], "'" + text)
        self.assertEqual(row["number"], "-7")
        write_output(store, self.root / "view.xlsx")
        book = load_workbook(self.root / "view.xlsx")
        try:
            headings = [cell.value for cell in book.active[1]]
            subject = headings.index("subject") + 1
            number = headings.index("number") + 1
            self.assertEqual(book.active.cell(2, subject).value, text)
            self.assertEqual(book.active.cell(2, subject).data_type, "s")
            self.assertEqual(book.active.cell(2, number).value, -7)
        finally:
            book.close()
